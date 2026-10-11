"""
FTS5 全文检索服务。

索引范围：规则、素材、项目/集/场景文本、技能插件文档。
"""
from __future__ import annotations

import json
import logging
import os
import shutil
import subprocess
import tempfile
from datetime import datetime
from pathlib import Path
from typing import Any

from sqlalchemy import text
from sqlalchemy.orm import Session

from app.models.material_item import MaterialItem
from app.models.project_v2 import Episode, ProjectV2, Scene
from app.models.rule import Rule


FTS_TABLE = "search_documents_fts"
ALLOWED_DOC_TYPES = {"rule", "material", "project", "skill"}
DEFAULT_LIMIT = 20
MAX_LIMIT = 100
PROJECT_ROOT = Path(__file__).resolve().parents[3]
RULE_MARKDOWN_ROOT = PROJECT_ROOT / "02-creation-rules"
MATERIAL_LIBRARY_ROOT = PROJECT_ROOT / "04-material-library"
SKILL_LIBRARY_ROOT = PROJECT_ROOT / "05-skill-plugins"
RULE_FILE_EXTENSIONS = {".md", ".markdown"}
TEXT_FILE_EXTENSIONS = {".md", ".markdown", ".txt", ".json", ".jsonl"}
IMAGE_FILE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}
MATERIAL_FILE_EXTENSIONS = TEXT_FILE_EXTENSIONS | {".docx", ".pdf"} | IMAGE_FILE_EXTENSIONS
SKILL_FILE_EXTENSIONS = {".md", ".markdown", ".json"}
SKILL_LIBRARY_EXCLUDED_PARTS = {
    "__pycache__",
    "node_modules",
    "dist",
    "build",
    "third_party",
    ".git",
}
MAX_INDEXABLE_FILE_BYTES = 2 * 1024 * 1024
MAX_EXTRACTED_FILE_CHARS = 120_000
MAX_PDF_PAGES = 80
MAX_OCR_IMAGE_PIXELS = 18_000_000
MAX_OCR_PDF_PAGES = 8
MATERIAL_OCR_ENV = "MOXI_ENABLE_MATERIAL_OCR"
MATERIAL_OCR_LANGS_ENV = "MOXI_MATERIAL_OCR_LANGS"

logger = logging.getLogger(__name__)


class FullTextSearchUnavailable(RuntimeError):
    """当前 SQLite 运行时不可用 FTS5。"""


def ensure_fts_index(db: Session) -> bool:
    """确保 FTS5 虚拟表存在；trigram 不可用时退回 unicode61。"""
    ddl_variants = (
        f"""
        CREATE VIRTUAL TABLE IF NOT EXISTS {FTS_TABLE}
        USING fts5(
            doc_type UNINDEXED,
            source_id UNINDEXED,
            title,
            content,
            metadata_json UNINDEXED,
            updated_at UNINDEXED,
            tokenize='trigram'
        )
        """,
        f"""
        CREATE VIRTUAL TABLE IF NOT EXISTS {FTS_TABLE}
        USING fts5(
            doc_type UNINDEXED,
            source_id UNINDEXED,
            title,
            content,
            metadata_json UNINDEXED,
            updated_at UNINDEXED
        )
        """,
    )
    last_error: Exception | None = None
    for ddl in ddl_variants:
        try:
            db.execute(text(ddl))
            return True
        except Exception as exc:
            last_error = exc
            db.rollback()
    raise FullTextSearchUnavailable(str(last_error or "FTS5 unavailable"))


def rebuild_search_index(db: Session) -> dict[str, Any]:
    """重建全文索引。"""
    try:
        ensure_fts_index(db)
        db.execute(text(f"DELETE FROM {FTS_TABLE}"))
        documents = list(_iter_search_documents(db))
        if documents:
            db.execute(
                text(
                    f"""
                    INSERT INTO {FTS_TABLE}
                        (doc_type, source_id, title, content, metadata_json, updated_at)
                    VALUES
                        (:doc_type, :source_id, :title, :content, :metadata_json, :updated_at)
                    """
                ),
                documents,
            )
        db.commit()
        return {
            "available": True,
            "indexed_count": len(documents),
            "doc_type_counts": _count_by_type(documents),
        }
    except FullTextSearchUnavailable:
        db.rollback()
        return {"available": False, "indexed_count": 0, "doc_type_counts": {}}
    except Exception:
        db.rollback()
        raise


def search_documents(
    db: Session,
    query: str,
    doc_types: list[str] | None = None,
    limit: int = DEFAULT_LIMIT,
    offset: int = 0,
    auto_rebuild: bool = True,
) -> dict[str, Any]:
    """检索全文索引；必要时自动重建，FTS 查询失败时退回 LIKE。"""
    normalized_query = (query or "").strip()
    if not normalized_query:
        return {
            "items": [],
            "count": 0,
            "indexed_count": _safe_indexed_count(db),
            "backend": "empty_query",
            "rebuilt": False,
        }

    limit = _clamp_limit(limit)
    offset = max(0, int(offset or 0))
    normalized_types = _normalize_doc_types(doc_types)

    try:
        ensure_fts_index(db)
    except FullTextSearchUnavailable:
        return _search_sources_like(db, normalized_query, normalized_types, limit, offset)

    rebuilt = False
    indexed_count = _safe_indexed_count(db)
    if auto_rebuild and indexed_count == 0:
        rebuild_result = rebuild_search_index(db)
        rebuilt = bool(rebuild_result.get("available"))
        indexed_count = int(rebuild_result.get("indexed_count") or 0)

    if _should_use_like(normalized_query):
        like_result = _search_index_like(db, normalized_query, normalized_types, limit, offset)
        like_result["indexed_count"] = indexed_count
        like_result["rebuilt"] = rebuilt
        return like_result

    try:
        fts_items = _search_index_fts(db, normalized_query, normalized_types, limit, offset)
    except Exception:
        db.rollback()
        fts_items = []

    if fts_items:
        return {
            "items": fts_items,
            "count": len(fts_items),
            "indexed_count": indexed_count,
            "backend": "fts5",
            "rebuilt": rebuilt,
        }

    like_result = _search_index_like(db, normalized_query, normalized_types, limit, offset)
    like_result["indexed_count"] = indexed_count
    like_result["rebuilt"] = rebuilt
    return like_result


def get_search_status(db: Session) -> dict[str, Any]:
    """返回当前全文索引状态。"""
    try:
        ensure_fts_index(db)
        indexed_count = _safe_indexed_count(db)
        rows = db.execute(
            text(f"SELECT doc_type, count(*) AS count FROM {FTS_TABLE} GROUP BY doc_type")
        ).mappings().all()
        return {
            "available": True,
            "indexed_count": indexed_count,
            "doc_type_counts": {row["doc_type"]: row["count"] for row in rows},
        }
    except FullTextSearchUnavailable:
        return {"available": False, "indexed_count": 0, "doc_type_counts": {}}


def get_material_ingestion_audit(limit: int = 200) -> dict[str, Any]:
    """扫描素材库摄取边界，不执行 OCR，不写库。"""
    limit = max(1, min(int(limit or 200), 1000))
    supported_extensions = sorted(MATERIAL_FILE_EXTENSIONS)
    summary = {
        "available": MATERIAL_LIBRARY_ROOT.exists() and MATERIAL_LIBRARY_ROOT.is_dir(),
        "root": str(MATERIAL_LIBRARY_ROOT),
        "ocr_enabled": _material_ocr_enabled(),
        "ocr_env": MATERIAL_OCR_ENV,
        "ocr_langs": _material_ocr_languages(),
        "supported_extensions": supported_extensions,
        "max_indexable_file_bytes": MAX_INDEXABLE_FILE_BYTES,
        "max_extracted_file_chars": MAX_EXTRACTED_FILE_CHARS,
        "max_pdf_pages": MAX_PDF_PAGES,
        "max_ocr_pdf_pages": MAX_OCR_PDF_PAGES,
        "max_ocr_image_pixels": MAX_OCR_IMAGE_PIXELS,
        "total_files": 0,
        "supported_files": 0,
        "oversized_files": 0,
        "ocr_candidate_files": 0,
        "review_required_files": 0,
        "status_counts": {},
        "extension_counts": {},
        "samples": [],
    }
    if not summary["available"]:
        return summary

    try:
        paths = sorted(MATERIAL_LIBRARY_ROOT.rglob("*"), key=lambda item: item.as_posix())
    except OSError as exc:
        logger.warning("skip unreadable material audit root: %s (%s)", MATERIAL_LIBRARY_ROOT, exc)
        summary["available"] = False
        summary["error"] = str(exc)
        return summary

    for path in paths:
        if not path.is_file():
            continue
        summary["total_files"] += 1
        extension = path.suffix.lower()
        _increment(summary["extension_counts"], extension or "<none>")
        if extension not in MATERIAL_FILE_EXTENSIONS:
            _increment(summary["status_counts"], "unsupported_extension")
            continue

        summary["supported_files"] += 1
        status = _material_ingestion_status(path, extension)
        _increment(summary["status_counts"], status["status"])
        if status["oversized"]:
            summary["oversized_files"] += 1
        if status["ocr_candidate"]:
            summary["ocr_candidate_files"] += 1
        if status["requires_review"]:
            summary["review_required_files"] += 1
        if len(summary["samples"]) < limit:
            try:
                relative_path = path.relative_to(MATERIAL_LIBRARY_ROOT).as_posix()
            except ValueError:
                relative_path = path.as_posix()
            summary["samples"].append({
                "path": relative_path,
                "extension": extension,
                **status,
            })

    return summary


def iter_indexable_documents(db: Session) -> list[dict[str, str]]:
    """返回可供 FTS5/Chroma 等检索后端复用的统一文档集合。"""
    return list(_iter_search_documents(db))


def _iter_search_documents(db: Session):
    for rule in db.query(Rule).order_by(Rule.updated_at.desc(), Rule.created_at.desc()).all():
        yield _rule_document(rule)

    for item in db.query(MaterialItem).order_by(MaterialItem.updated_at.desc(), MaterialItem.created_at.desc()).all():
        yield _material_document(item)

    projects = db.query(ProjectV2).order_by(ProjectV2.updated_at.desc()).all()
    projects_by_id = {project.id: project for project in projects}
    for project in projects:
        yield _project_document(project)

    episodes = db.query(Episode).order_by(Episode.updated_at.desc()).all()
    episodes_by_id = {episode.id: episode for episode in episodes}
    for episode in episodes:
        project = projects_by_id.get(episode.project_id)
        yield _episode_document(episode, project)

    for scene in db.query(Scene).order_by(Scene.updated_at.desc()).all():
        episode = episodes_by_id.get(scene.episode_id)
        if not episode:
            continue
        project = projects_by_id.get(episode.project_id)
        yield _scene_document(scene, episode, project)

    yield from _iter_file_library_documents()


def _rule_document(rule: Rule) -> dict[str, str]:
    return _document(
        doc_type="rule",
        source_id=rule.id,
        title=f"规则 {rule.scope} / {rule.rule_type} / {rule.status}",
        content=rule.content,
        metadata={
            "kind": "rule",
            "scope": rule.scope,
            "task_id": rule.task_id,
            "sub_task_id": rule.sub_task_id,
            "rule_type": rule.rule_type,
            "status": rule.status,
            "source": rule.source,
        },
        updated_at=rule.updated_at or rule.created_at,
    )


def _material_document(item: MaterialItem) -> dict[str, str]:
    project_tags = _json_loads(item.project_tags_json, [])
    data_points = _json_loads(item.data_points_json, [])
    content_parts = [
        item.summary,
        item.heat_reason,
        " ".join(str(tag) for tag in project_tags),
        " ".join(str(point) for point in data_points),
    ]
    return _document(
        doc_type="material",
        source_id=item.id,
        title=item.title,
        content="\n".join(part for part in content_parts if part),
        metadata={
            "kind": "material",
            "source_platform": item.source_platform,
            "source_url": item.source_url,
            "content_type": item.content_type,
            "heat_score": item.heat_score,
            "selection_score": item.selection_score,
            "status": item.status,
            "project_tags": project_tags,
        },
        updated_at=item.updated_at or item.created_at,
    )


def _project_document(project: ProjectV2) -> dict[str, str]:
    content_parts = [
        project.description,
        project.project_type,
        project.audience_axis,
        project.genre_main,
        project.genre_sub,
        json.dumps(project.get_config(), ensure_ascii=False),
    ]
    return _document(
        doc_type="project",
        source_id=f"project:{project.id}",
        title=project.name,
        content="\n".join(part for part in content_parts if part),
        metadata={
            "kind": "project",
            "project_id": project.id,
            "status": project.status,
            "workspace_path": project.workspace_path,
        },
        updated_at=project.updated_at or project.created_at,
    )


def _episode_document(episode: Episode, project: ProjectV2 | None) -> dict[str, str]:
    project_name = project.name if project else episode.project_id
    return _document(
        doc_type="project",
        source_id=f"episode:{episode.id}",
        title=f"{project_name} 第{episode.episode_number}集 {episode.title}".strip(),
        content=episode.summary,
        metadata={
            "kind": "episode",
            "project_id": episode.project_id,
            "episode_id": episode.id,
            "episode_number": episode.episode_number,
            "status": episode.status,
        },
        updated_at=episode.updated_at or episode.created_at,
    )


def _scene_document(scene: Scene, episode: Episode, project: ProjectV2 | None) -> dict[str, str]:
    project_name = project.name if project else episode.project_id
    title = f"{project_name} 第{episode.episode_number}集 第{scene.scene_number}场 {scene.location_name}".strip()
    content_parts = [
        scene.scene_code,
        scene.time,
        scene.location_type,
        scene.location_name,
        scene.content,
    ]
    return _document(
        doc_type="project",
        source_id=f"scene:{scene.id}",
        title=title,
        content="\n".join(part for part in content_parts if part),
        metadata={
            "kind": "scene",
            "project_id": episode.project_id,
            "episode_id": episode.id,
            "scene_id": scene.id,
            "episode_number": episode.episode_number,
            "scene_number": scene.scene_number,
            "status": scene.status,
        },
        updated_at=scene.updated_at or scene.created_at,
    )


def _iter_file_library_documents():
    yield from _iter_files_under_root(
        root=RULE_MARKDOWN_ROOT,
        doc_type="rule",
        kind="rule_file",
        library="02-creation-rules",
        extensions=RULE_FILE_EXTENSIONS,
    )
    yield from _iter_files_under_root(
        root=MATERIAL_LIBRARY_ROOT,
        doc_type="material",
        kind="material_file",
        library="04-material-library",
        extensions=MATERIAL_FILE_EXTENSIONS,
    )
    yield from _iter_files_under_root(
        root=SKILL_LIBRARY_ROOT,
        doc_type="skill",
        kind="skill_file",
        library="05-skill-plugins",
        extensions=SKILL_FILE_EXTENSIONS,
        excluded_parts=SKILL_LIBRARY_EXCLUDED_PARTS,
    )


def _iter_files_under_root(
    *,
    root: Path,
    doc_type: str,
    kind: str,
    library: str,
    extensions: set[str],
    excluded_parts: set[str] | None = None,
):
    if not root.exists() or not root.is_dir():
        return

    try:
        paths = sorted(root.rglob("*"), key=lambda item: item.as_posix())
    except OSError as exc:
        logger.warning("skip unreadable search library root: %s (%s)", root, exc)
        return

    for path in paths:
        if not _is_indexable_file_path(path, root, extensions, excluded_parts or set()):
            continue
        document = _file_library_document(
            path=path,
            root=root,
            doc_type=doc_type,
            kind=kind,
            library=library,
        )
        if document:
            yield document


def _is_indexable_file_path(
    path: Path,
    root: Path,
    extensions: set[str],
    excluded_parts: set[str],
) -> bool:
    if not path.is_file() or path.suffix.lower() not in extensions:
        return False
    try:
        relative_parts = path.relative_to(root).parts
    except ValueError:
        return False
    return not any(part.startswith(".") or part in excluded_parts for part in relative_parts)


def _file_library_document(
    *,
    path: Path,
    root: Path,
    doc_type: str,
    kind: str,
    library: str,
) -> dict[str, str] | None:
    try:
        stat = path.stat()
    except OSError as exc:
        logger.warning("skip unreadable search file stat: %s (%s)", path, exc)
        return None

    if stat.st_size > MAX_INDEXABLE_FILE_BYTES:
        logger.warning("skip oversized search file: %s (%s bytes)", path, stat.st_size)
        return None

    extraction = _extract_indexable_file_text(path)
    if extraction is None:
        return None

    relative_path = path.relative_to(root).as_posix()
    normalized_content = extraction["content"]
    return _document(
        doc_type=doc_type,
        source_id=f"file:{library}:{relative_path}",
        title=path.stem,
        content=normalized_content,
        metadata={
            "kind": kind,
            "library": library,
            "path": relative_path,
            "extension": path.suffix.lower(),
            "size_bytes": stat.st_size,
            "parser": extraction.get("parser", ""),
            "content_chars": len(normalized_content),
            "truncated": bool(extraction.get("truncated")),
            **_material_governance_metadata(
                library=library,
                extension=path.suffix.lower(),
                parser=str(extraction.get("parser") or ""),
                extraction=extraction,
            ),
            **_as_dict(extraction.get("metadata")),
        },
        updated_at=datetime.fromtimestamp(stat.st_mtime),
    )


def _extract_indexable_file_text(path: Path) -> dict[str, Any] | None:
    extension = path.suffix.lower()
    if extension in TEXT_FILE_EXTENSIONS:
        content = _read_text_file(path)
        if content is None:
            return None
        normalized = _normalize_file_content(content, extension)
        normalized, truncated = _truncate_extracted_text(normalized)
        return {
            "content": normalized,
            "parser": "utf8_text",
            "truncated": truncated,
            "metadata": {},
        }
    if extension == ".docx":
        return _extract_docx_text(path)
    if extension == ".pdf":
        return _extract_pdf_text(path)
    if extension in IMAGE_FILE_EXTENSIONS:
        return _extract_image_ocr_text(path)
    return None


def _read_text_file(path: Path) -> str | None:
    try:
        return path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        try:
            return path.read_text(encoding="utf-8-sig")
        except Exception as exc:
            logger.warning("skip undecodable search file: %s (%s)", path, exc)
            return None
    except OSError as exc:
        logger.warning("skip unreadable search file: %s (%s)", path, exc)
        return None


def _extract_docx_text(path: Path) -> dict[str, Any] | None:
    try:
        from docx import Document
    except Exception as exc:
        logger.warning("skip docx search file without python-docx: %s (%s)", path, exc)
        return None

    try:
        document = Document(str(path))
    except Exception as exc:
        logger.warning("skip unreadable docx search file: %s (%s)", path, exc)
        return None

    parts: list[str] = []
    truncated = False
    for paragraph in document.paragraphs:
        if not _append_bounded_text(parts, paragraph.text):
            truncated = True
            break
    if not truncated:
        for table in document.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text and not _append_bounded_text(parts, row_text):
                    truncated = True
                    break
            if truncated:
                break
    content = "\n".join(parts).strip()
    if not content:
        logger.warning("skip empty docx search file: %s", path)
        return None
    return {
        "content": content,
        "parser": "python-docx",
        "truncated": truncated,
        "metadata": {
            "paragraph_count": len(document.paragraphs),
            "table_count": len(document.tables),
        },
    }


def _extract_pdf_text(path: Path) -> dict[str, Any] | None:
    try:
        from pypdf import PdfReader
    except Exception as exc:
        logger.warning("skip pdf search file without pypdf: %s (%s)", path, exc)
        return None

    try:
        reader = PdfReader(str(path))
    except Exception as exc:
        logger.warning("skip unreadable pdf search file: %s (%s)", path, exc)
        return None
    if getattr(reader, "is_encrypted", False):
        logger.warning("skip encrypted pdf search file: %s", path)
        return None

    parts: list[str] = []
    pages = list(reader.pages)
    truncated = len(pages) > MAX_PDF_PAGES
    pages_indexed = 0
    for page in pages[:MAX_PDF_PAGES]:
        pages_indexed += 1
        try:
            page_text = page.extract_text() or ""
        except Exception as exc:
            logger.warning("skip unreadable pdf page: %s page=%s (%s)", path, pages_indexed, exc)
            continue
        if page_text and not _append_bounded_text(parts, page_text):
            truncated = True
            break
    content = "\n".join(parts).strip()
    if not content:
        if _material_ocr_enabled():
            return _extract_pdf_ocr_text(path, page_count=len(pages))
        logger.warning("skip empty pdf search file without OCR enabled: %s", path)
        return None
    return {
        "content": content,
        "parser": "pypdf",
        "truncated": truncated,
        "metadata": {
            "page_count": len(pages),
            "pages_indexed": pages_indexed,
        },
    }


def _extract_image_ocr_text(path: Path) -> dict[str, Any] | None:
    if not _material_ocr_enabled():
        logger.warning("skip image material OCR because %s is not enabled: %s", MATERIAL_OCR_ENV, path)
        return None

    try:
        from PIL import Image
    except Exception as exc:
        logger.warning("skip image OCR without Pillow: %s (%s)", path, exc)
        return None

    try:
        with Image.open(path) as image:
            image.load()
            width, height = image.size
            if width * height > MAX_OCR_IMAGE_PIXELS:
                logger.warning("skip oversized image OCR: %s (%sx%s)", path, width, height)
                return None
            text_value = _extract_text_from_image(image)
    except Exception as exc:
        logger.warning("skip unreadable image OCR file: %s (%s)", path, exc)
        return None

    content, truncated = _truncate_extracted_text(text_value.strip())
    if not content:
        logger.warning("skip empty image OCR file: %s", path)
        return None
    return {
        "content": content,
        "parser": "tesseract_image",
        "truncated": truncated,
        "metadata": {
            "ocr_used": True,
            "ocr_source": "image",
            "ocr_langs": _material_ocr_languages(),
            "ocr_engine": "tesseract",
        },
    }


def _extract_pdf_ocr_text(path: Path, page_count: int) -> dict[str, Any] | None:
    pypdfium_result = _extract_pdf_ocr_text_with_pypdfium2(path, page_count)
    if pypdfium_result is not None:
        return pypdfium_result
    return _extract_pdf_ocr_text_with_pdftoppm(path, page_count)


def _extract_pdf_ocr_text_with_pypdfium2(path: Path, page_count: int) -> dict[str, Any] | None:
    try:
        import pypdfium2 as pdfium
    except Exception as exc:
        logger.debug("pypdfium2 unavailable for scanned pdf OCR, try pdftoppm fallback: %s (%s)", path, exc)
        return None

    try:
        pdf = pdfium.PdfDocument(str(path))
    except Exception as exc:
        logger.warning("skip unreadable scanned pdf OCR file: %s (%s)", path, exc)
        return None

    parts: list[str] = []
    truncated = page_count > MAX_OCR_PDF_PAGES
    pages_indexed = 0
    try:
        for index in range(min(len(pdf), MAX_OCR_PDF_PAGES)):
            pages_indexed += 1
            page = pdf[index]
            try:
                bitmap = page.render(scale=2)
                image = bitmap.to_pil()
                page_text = _extract_text_from_image(image)
            except Exception as exc:
                logger.warning("skip unreadable scanned pdf OCR page: %s page=%s (%s)", path, pages_indexed, exc)
                continue
            if page_text and not _append_bounded_text(parts, page_text):
                truncated = True
                break
    finally:
        close = getattr(pdf, "close", None)
        if callable(close):
            close()

    content = "\n".join(parts).strip()
    if not content:
        logger.warning("skip empty scanned pdf OCR file: %s", path)
        return None
    return {
        "content": content,
        "parser": "tesseract_pdf",
        "truncated": truncated,
        "metadata": {
            "page_count": page_count,
            "pages_indexed": pages_indexed,
            "ocr_used": True,
            "ocr_source": "scanned_pdf",
            "ocr_langs": _material_ocr_languages(),
            "ocr_engine": "tesseract",
        },
    }


def _extract_pdf_ocr_text_with_pdftoppm(path: Path, page_count: int) -> dict[str, Any] | None:
    if not shutil.which("pdftoppm"):
        logger.warning("skip scanned pdf OCR without pdftoppm: %s", path)
        return None

    pages_to_render = min(page_count, MAX_OCR_PDF_PAGES)
    parts: list[str] = []
    truncated = page_count > MAX_OCR_PDF_PAGES
    pages_indexed = 0
    with tempfile.TemporaryDirectory(prefix="moxi-pdf-ocr-") as tmpdir:
        prefix = str(Path(tmpdir) / "page")
        cmd = [
            "pdftoppm",
            "-png",
            "-r",
            "180",
            "-f",
            "1",
            "-l",
            str(pages_to_render),
            str(path),
            prefix,
        ]
        try:
            subprocess.run(cmd, capture_output=True, text=True, check=True, timeout=90)
        except Exception as exc:
            logger.warning("skip scanned pdf OCR render failure: %s (%s)", path, exc)
            return None

        for image_path in sorted(Path(tmpdir).glob("page-*.png")):
            pages_indexed += 1
            try:
                page_text = _extract_text_from_image_path(image_path)
            except Exception as exc:
                logger.warning("skip scanned pdf OCR rendered page: %s page=%s (%s)", path, pages_indexed, exc)
                continue
            if page_text and not _append_bounded_text(parts, page_text):
                truncated = True
                break

    content = "\n".join(parts).strip()
    if not content:
        logger.warning("skip empty scanned pdf OCR file: %s", path)
        return None
    return {
        "content": content,
        "parser": "tesseract_pdf",
        "truncated": truncated,
        "metadata": {
            "page_count": page_count,
            "pages_indexed": pages_indexed,
            "ocr_used": True,
            "ocr_source": "scanned_pdf",
            "ocr_langs": _material_ocr_languages(),
            "ocr_engine": "tesseract",
            "pdf_renderer": "pdftoppm",
        },
    }


def _extract_text_from_image(image: Any) -> str:
    try:
        import pytesseract
    except Exception as exc:
        return _extract_text_from_image_cli(image)

    langs = _material_ocr_languages()
    try:
        return str(pytesseract.image_to_string(image, lang=langs) or "")
    except Exception:
        if langs != "eng":
            return str(pytesseract.image_to_string(image, lang="eng") or "")
        raise


def _extract_text_from_image_path(path: Path) -> str:
    try:
        from PIL import Image
    except Exception as exc:
        raise RuntimeError(f"Pillow unavailable: {exc}") from exc
    with Image.open(path) as image:
        image.load()
        return _extract_text_from_image(image)


def _extract_text_from_image_cli(image: Any) -> str:
    if not shutil.which("tesseract"):
        raise RuntimeError("tesseract executable unavailable")
    with tempfile.NamedTemporaryFile(suffix=".png") as tmp:
        image.save(tmp.name)
        return _run_tesseract_cli(tmp.name, _material_ocr_languages())


def _run_tesseract_cli(image_path: str, langs: str) -> str:
    cmd = ["tesseract", image_path, "stdout", "-l", langs]
    try:
        completed = subprocess.run(cmd, capture_output=True, text=True, check=True, timeout=60)
        return completed.stdout or ""
    except Exception:
        if langs != "eng":
            completed = subprocess.run(
                ["tesseract", image_path, "stdout", "-l", "eng"],
                capture_output=True,
                text=True,
                check=True,
                timeout=60,
            )
            return completed.stdout or ""
        raise


def _append_bounded_text(parts: list[str], text_value: str) -> bool:
    text = str(text_value or "").strip()
    if not text:
        return True
    current_chars = sum(len(part) + 1 for part in parts)
    remaining = MAX_EXTRACTED_FILE_CHARS - current_chars
    if remaining <= 0:
        return False
    if len(text) > remaining:
        parts.append(text[:remaining])
        return False
    parts.append(text)
    return True


def _truncate_extracted_text(content: str) -> tuple[str, bool]:
    if len(content) <= MAX_EXTRACTED_FILE_CHARS:
        return content, False
    return content[:MAX_EXTRACTED_FILE_CHARS], True


def _normalize_file_content(content: str, extension: str) -> str:
    if extension != ".json":
        return content
    try:
        return json.dumps(json.loads(content), ensure_ascii=False, indent=2)
    except Exception:
        return content


def _document(
    *,
    doc_type: str,
    source_id: str,
    title: str,
    content: str,
    metadata: dict[str, Any],
    updated_at: datetime | None,
) -> dict[str, str]:
    return {
        "doc_type": doc_type,
        "source_id": source_id,
        "title": str(title or ""),
        "content": str(content or ""),
        "metadata_json": json.dumps(metadata, ensure_ascii=False, default=str),
        "updated_at": _dt_to_str(updated_at),
    }


def _search_index_fts(
    db: Session,
    query: str,
    doc_types: list[str],
    limit: int,
    offset: int,
) -> list[dict[str, Any]]:
    where_parts = [f"{FTS_TABLE} MATCH :match_query"]
    params: dict[str, Any] = {
        "match_query": _fts_match_query(query),
        "limit": limit,
        "offset": offset,
    }
    _apply_type_filter(where_parts, params, doc_types)
    rows = db.execute(
        text(
            f"""
            SELECT
                rowid,
                doc_type,
                source_id,
                title,
                snippet({FTS_TABLE}, 2, '<mark>', '</mark>', '...', 12) AS title_snippet,
                snippet({FTS_TABLE}, 3, '<mark>', '</mark>', '...', 24) AS content_snippet,
                metadata_json,
                updated_at,
                bm25({FTS_TABLE}) AS score
            FROM {FTS_TABLE}
            WHERE {" AND ".join(where_parts)}
            ORDER BY score ASC, rowid DESC
            LIMIT :limit OFFSET :offset
            """
        ),
        params,
    ).mappings().all()
    return [_row_to_result(row, "fts5") for row in rows]


def _search_index_like(
    db: Session,
    query: str,
    doc_types: list[str],
    limit: int,
    offset: int,
) -> dict[str, Any]:
    where_parts = ["(title LIKE :pattern OR content LIKE :pattern)"]
    params: dict[str, Any] = {
        "pattern": f"%{query}%",
        "limit": limit,
        "offset": offset,
    }
    _apply_type_filter(where_parts, params, doc_types)
    rows = db.execute(
        text(
            f"""
            SELECT
                rowid,
                doc_type,
                source_id,
                title,
                title AS title_snippet,
                content AS content_snippet,
                metadata_json,
                updated_at,
                0.0 AS score
            FROM {FTS_TABLE}
            WHERE {" AND ".join(where_parts)}
            ORDER BY updated_at DESC, rowid DESC
            LIMIT :limit OFFSET :offset
            """
        ),
        params,
    ).mappings().all()
    items = [_row_to_result(row, "like_fallback") for row in rows]
    return {
        "items": items,
        "count": len(items),
        "indexed_count": _safe_indexed_count(db),
        "backend": "like_fallback",
        "rebuilt": False,
    }


def _search_sources_like(
    db: Session,
    query: str,
    doc_types: list[str],
    limit: int,
    offset: int,
) -> dict[str, Any]:
    documents = [
        doc
        for doc in _iter_search_documents(db)
        if (not doc_types or doc["doc_type"] in doc_types)
        and (query in doc["title"] or query in doc["content"])
    ]
    page = documents[offset: offset + limit]
    items = [
        {
            "doc_type": doc["doc_type"],
            "source_id": doc["source_id"],
            "title": doc["title"],
            "title_snippet": doc["title"],
            "content_snippet": _plain_snippet(doc["content"], query),
            "metadata": _json_loads(doc["metadata_json"], {}),
            "updated_at": doc["updated_at"],
            "score": 0.0,
            "backend": "source_like_fallback",
        }
        for doc in page
    ]
    return {
        "items": items,
        "count": len(items),
        "indexed_count": 0,
        "backend": "source_like_fallback",
        "rebuilt": False,
    }


def _apply_type_filter(where_parts: list[str], params: dict[str, Any], doc_types: list[str]) -> None:
    if not doc_types:
        return
    placeholders = []
    for index, doc_type in enumerate(doc_types):
        key = f"doc_type_{index}"
        params[key] = doc_type
        placeholders.append(f":{key}")
    where_parts.append(f"doc_type IN ({', '.join(placeholders)})")


def _row_to_result(row, backend: str) -> dict[str, Any]:
    return {
        "doc_type": row["doc_type"],
        "source_id": row["source_id"],
        "title": row["title"],
        "title_snippet": row["title_snippet"] or row["title"],
        "content_snippet": row["content_snippet"] or "",
        "metadata": _json_loads(row["metadata_json"], {}),
        "updated_at": row["updated_at"],
        "score": float(row["score"] or 0.0),
        "backend": backend,
    }


def _safe_indexed_count(db: Session) -> int:
    try:
        ensure_fts_index(db)
        return int(db.execute(text(f"SELECT count(*) FROM {FTS_TABLE}")).scalar() or 0)
    except Exception:
        db.rollback()
        return 0


def _normalize_doc_types(doc_types: list[str] | None) -> list[str]:
    normalized = []
    for item in doc_types or []:
        for part in str(item).split(","):
            doc_type = part.strip()
            if doc_type and doc_type in ALLOWED_DOC_TYPES and doc_type not in normalized:
                normalized.append(doc_type)
    return normalized


def _clamp_limit(limit: int) -> int:
    try:
        value = int(limit)
    except Exception:
        value = DEFAULT_LIMIT
    return max(1, min(value, MAX_LIMIT))


def _should_use_like(query: str) -> bool:
    return len(query.strip()) < 3


def _fts_match_query(query: str) -> str:
    escaped = query.strip().replace('"', '""')
    return f'"{escaped}"'


def _count_by_type(documents: list[dict[str, str]]) -> dict[str, int]:
    counts: dict[str, int] = {}
    for doc in documents:
        doc_type = doc["doc_type"]
        counts[doc_type] = counts.get(doc_type, 0) + 1
    return counts


def _json_loads(raw: str | None, default: Any) -> Any:
    if not raw:
        return default
    try:
        return json.loads(raw)
    except Exception:
        return default


def _as_dict(value: Any) -> dict[str, Any]:
    return value if isinstance(value, dict) else {}


def _material_ocr_enabled() -> bool:
    return str(os.getenv(MATERIAL_OCR_ENV, "")).strip().lower() in {"1", "true", "yes", "on"}


def _material_ocr_languages() -> str:
    return str(os.getenv(MATERIAL_OCR_LANGS_ENV, "chi_sim+eng")).strip() or "eng"


def _material_governance_metadata(
    *,
    library: str,
    extension: str,
    parser: str,
    extraction: dict[str, Any],
) -> dict[str, Any]:
    if library != "04-material-library":
        return {}
    ocr_used = parser.startswith("tesseract_") or bool(_as_dict(extraction.get("metadata")).get("ocr_used"))
    requires_review = ocr_used or extension in IMAGE_FILE_EXTENSIONS or extension == ".pdf"
    flags: list[str] = ["local_material_library", "reference_only"]
    if ocr_used:
        flags.extend(["ocr_extracted", "copyright_review_required"])
    elif extension == ".pdf":
        flags.extend(["pdf_material", "copyright_review_required"])
    elif extension in IMAGE_FILE_EXTENSIONS:
        flags.extend(["image_material", "copyright_review_required"])
    return {
        "material_ingestion_policy": "reference_only",
        "requires_rights_review": requires_review,
        "allowed_for_direct_copy": False,
        "policy_flags": flags,
    }


def _material_ingestion_status(path: Path, extension: str) -> dict[str, Any]:
    try:
        stat = path.stat()
        size_bytes = stat.st_size
    except OSError:
        return {
            "status": "unreadable",
            "size_bytes": 0,
            "oversized": False,
            "ocr_candidate": False,
            "requires_review": False,
        }

    oversized = size_bytes > MAX_INDEXABLE_FILE_BYTES
    ocr_candidate = extension in IMAGE_FILE_EXTENSIONS
    requires_review = extension in IMAGE_FILE_EXTENSIONS or extension == ".pdf"
    if extension == ".pdf":
        status = "pdf_text_or_ocr_candidate" if not oversized else "oversized"
        ocr_candidate = True
    elif extension in IMAGE_FILE_EXTENSIONS:
        status = "image_ocr_enabled" if _material_ocr_enabled() else "image_ocr_disabled"
        if oversized:
            status = "oversized"
    elif extension in TEXT_FILE_EXTENSIONS or extension == ".docx":
        status = "indexable" if not oversized else "oversized"
    else:
        status = "unsupported_extension"
    return {
        "status": status,
        "size_bytes": size_bytes,
        "oversized": oversized,
        "ocr_candidate": ocr_candidate,
        "requires_review": requires_review,
    }


def _increment(counts: dict[str, int], key: str) -> None:
    counts[key] = int(counts.get(key) or 0) + 1


def _dt_to_str(value: datetime | None) -> str:
    if not value:
        return ""
    return value.isoformat(sep=" ", timespec="seconds")


def _plain_snippet(content: str, query: str, radius: int = 36) -> str:
    if not content:
        return ""
    index = content.find(query)
    if index < 0:
        return content[: radius * 2]
    start = max(0, index - radius)
    end = min(len(content), index + len(query) + radius)
    prefix = "..." if start > 0 else ""
    suffix = "..." if end < len(content) else ""
    return f"{prefix}{content[start:end]}{suffix}"
