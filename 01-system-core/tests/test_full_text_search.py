from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool


def _isolate_file_libraries(monkeypatch, tmp_path):
    from app.services import full_text_search_service

    rule_root = tmp_path / "02-creation-rules"
    material_root = tmp_path / "04-material-library"
    skill_root = tmp_path / "05-skill-plugins"
    rule_root.mkdir()
    material_root.mkdir()
    skill_root.mkdir()
    monkeypatch.setattr(full_text_search_service, "RULE_MARKDOWN_ROOT", rule_root)
    monkeypatch.setattr(full_text_search_service, "MATERIAL_LIBRARY_ROOT", material_root)
    monkeypatch.setattr(full_text_search_service, "SKILL_LIBRARY_ROOT", skill_root)
    return rule_root, material_root


def _make_search_session():
    from app.models.agent import Agent
    from app.models.material_item import MaterialItem
    from app.models.module import Module
    from app.models.project_v2 import Episode, ProjectV2, Scene
    from app.models.rule import Rule
    from app.models.sub_task import SubTask
    from app.models.task import Task

    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    for table in (
        Agent.__table__,
        Task.__table__,
        Module.__table__,
        SubTask.__table__,
        Rule.__table__,
        MaterialItem.__table__,
        ProjectV2.__table__,
        Episode.__table__,
        Scene.__table__,
    ):
        table.create(engine)
    return sessionmaker(bind=engine)()


def test_full_text_search_rebuild_indexes_rules_materials_and_project_text(monkeypatch, tmp_path):
    from app.models.material_item import MaterialItem
    from app.models.project_v2 import Episode, ProjectV2, Scene
    from app.models.rule import Rule
    from app.services import full_text_search_service

    _isolate_file_libraries(monkeypatch, tmp_path)
    db = _make_search_session()
    try:
        db.add(Rule(
            id="rule-1",
            scope="global",
            rule_type="core_rule",
            status="active",
            source="manual",
            content="都市家庭短剧规则：人物要有真实生活压力。",
        ))
        db.add(MaterialItem(
            id="material-1",
            source_platform="红果",
            source_url="https://example.com/hot",
            content_type="短剧",
            title="外卖骑手爆款结构",
            summary="外卖骑手发现旧城秘密后反转逆袭。",
            heat_reason="旧城、家庭压力、现实创业具备热度。",
            project_tags_json='["都市", "现实创业"]',
            status="approved",
        ))
        db.add(ProjectV2(
            id="project-1",
            name="旧城早餐计划",
            description="一个外卖骑手守住旧城早餐铺的故事。",
            project_type="真人短剧-男频",
            genre_main="现实创业",
            workspace_path="/tmp/project-1",
        ))
        db.add(Episode(
            id="episode-1",
            project_id="project-1",
            episode_number=1,
            title="第一笼出锅",
            summary="主角发现连锁早餐品牌压价。",
        ))
        db.add(Scene(
            id="scene-1",
            episode_id="episode-1",
            scene_number=1,
            scene_code="1-1",
            location_name="旧城巷口",
            content="外卖骑手陆远发现旧城秘密，决定守住早餐铺。",
        ))
        db.commit()

        rebuild = full_text_search_service.rebuild_search_index(db)
        assert rebuild["available"] is True
        assert rebuild["indexed_count"] == 5
        assert rebuild["doc_type_counts"] == {"rule": 1, "material": 1, "project": 3}

        result = full_text_search_service.search_documents(db, "外卖骑手", limit=10)
        assert result["backend"] == "fts5"
        assert result["count"] >= 2
        assert {"material", "project"} <= {item["doc_type"] for item in result["items"]}
        assert any(item["metadata"].get("kind") == "scene" for item in result["items"])
    finally:
        db.close()


def test_full_text_search_filters_doc_types_and_falls_back_for_short_chinese_terms(monkeypatch, tmp_path):
    from app.models.material_item import MaterialItem
    from app.models.rule import Rule
    from app.services import full_text_search_service

    _isolate_file_libraries(monkeypatch, tmp_path)
    db = _make_search_session()
    try:
        db.add(Rule(
            id="rule-1",
            scope="global",
            rule_type="core_rule",
            status="active",
            source="manual",
            content="都市题材需要强冲突、强钩子、强人物压力。",
        ))
        db.add(MaterialItem(
            id="material-1",
            source_platform="红果",
            source_url="https://example.com/material",
            content_type="短剧",
            title="都市家庭热点",
            summary="家庭债务和亲情误会推动剧情。",
            status="approved",
        ))
        db.commit()

        full_text_search_service.rebuild_search_index(db)
        result = full_text_search_service.search_documents(
            db,
            "都市",
            doc_types=["rule"],
            limit=10,
        )

        assert result["backend"] == "like_fallback"
        assert result["count"] == 1
        assert result["items"][0]["doc_type"] == "rule"
        assert result["items"][0]["source_id"] == "rule-1"
    finally:
        db.close()


def test_search_router_is_registered_in_main_app():
    from app.main import app

    paths = set(app.openapi().get("paths", {}).keys())

    assert "/api/search" in paths
    assert "/api/search/status" in paths
    assert "/api/search/rebuild" in paths
    assert "/api/search/semantic" in paths
    assert "/api/search/semantic/status" in paths
    assert "/api/search/semantic/rebuild" in paths
    assert "/api/search/material-ingestion/audit" in paths
    assert "/api/search/chunks" in paths
    assert "/api/search/chunks/status" in paths
    assert "/api/search/chunks/context" in paths
    assert "/api/search/chunks/rebuild" in paths


def test_search_endpoint_uses_current_database_session_and_auth_override(monkeypatch, tmp_path):
    from fastapi.testclient import TestClient

    from app.auth.dependencies import verify_admin
    from app.database import get_db
    from app.main import app
    from app.models.project_v2 import Episode, ProjectV2, Scene

    _isolate_file_libraries(monkeypatch, tmp_path)
    db = _make_search_session()
    db.add(ProjectV2(
        id="project-1",
        name="旧城早餐计划",
        description="一个外卖骑手守住旧城早餐铺的故事。",
        project_type="真人短剧-男频",
    ))
    db.add(Episode(id="episode-1", project_id="project-1", episode_number=1, title="第一集"))
    db.add(Scene(
        id="scene-1",
        episode_id="episode-1",
        scene_number=1,
        location_name="旧城巷口",
        content="外卖骑手陆远发现旧城秘密。",
    ))
    db.commit()

    def override_db():
        yield db

    app.dependency_overrides[get_db] = override_db
    app.dependency_overrides[verify_admin] = lambda: True
    monkeypatch.setenv("MOXI_ENABLE_LOCAL_SEARCH", "1")
    try:
        client = TestClient(app)
        resp = client.get("/api/search", params={"q": "外卖骑手", "types": "project"})

        assert resp.status_code == 200, resp.text
        payload = resp.json()
        assert payload["count"] >= 1
        assert any(item["metadata"]["kind"] == "scene" for item in payload["items"])
    finally:
        app.dependency_overrides.pop(get_db, None)
        app.dependency_overrides.pop(verify_admin, None)
        db.close()


def test_search_endpoint_defaults_to_disabled(monkeypatch, tmp_path):
    from fastapi.testclient import TestClient

    from app.auth.dependencies import verify_admin
    from app.database import get_db
    from app.main import app

    monkeypatch.delenv("MOXI_ENABLE_LOCAL_SEARCH", raising=False)
    _isolate_file_libraries(monkeypatch, tmp_path)
    db = _make_search_session()

    def override_db():
        yield db

    app.dependency_overrides[get_db] = override_db
    app.dependency_overrides[verify_admin] = lambda: True
    try:
        client = TestClient(app)
        resp = client.get("/api/search", params={"q": "外卖骑手", "types": "project"})

        assert resp.status_code == 200, resp.text
        payload = resp.json()
        assert payload["backend"] == "local_search_disabled"
        assert payload["count"] == 0
    finally:
        app.dependency_overrides.pop(get_db, None)
        app.dependency_overrides.pop(verify_admin, None)
        db.close()


def test_full_text_search_indexes_rule_markdown_and_material_files(monkeypatch, tmp_path):
    from app.services import full_text_search_service

    rule_root, material_root = _isolate_file_libraries(monkeypatch, tmp_path)
    (rule_root / "core-rules").mkdir()
    (rule_root / "core-rules" / "hooks.md").write_text(
        "# 黄金三秒钩子\n\n开头必须有身份错位和强反转。",
        encoding="utf-8",
    )
    (material_root / "hot").mkdir()
    (material_root / "hot" / "case.json").write_text(
        '{"title": "下沉市场素材", "summary": "烟火气早餐摊与家庭债务冲突"}',
        encoding="utf-8",
    )
    (material_root / "hot" / "notes.txt").write_text(
        "外部素材文本：旧城烟火气和早餐摊高频出现。",
        encoding="utf-8",
    )
    (material_root / "hot" / "ignored.pdf").write_bytes(b"%PDF not indexed")

    db = _make_search_session()
    try:
        rebuild = full_text_search_service.rebuild_search_index(db)

        assert rebuild["available"] is True
        assert rebuild["indexed_count"] == 3
        assert rebuild["doc_type_counts"] == {"rule": 1, "material": 2}

        result = full_text_search_service.search_documents(db, "早餐摊", doc_types=["material"], limit=10)
        assert result["count"] == 2
        assert {item["metadata"]["kind"] for item in result["items"]} == {"material_file"}
        assert all(item["source_id"].startswith("file:04-material-library:") for item in result["items"])

        rule_result = full_text_search_service.search_documents(db, "身份错位", doc_types=["rule"], limit=10)
        assert rule_result["count"] == 1
        assert rule_result["items"][0]["metadata"]["kind"] == "rule_file"
    finally:
        db.close()


def test_full_text_search_indexes_docx_and_pdf_material_files(monkeypatch, tmp_path):
    from docx import Document

    from app.services import full_text_search_service

    _, material_root = _isolate_file_libraries(monkeypatch, tmp_path)
    (material_root / "scripts").mkdir()

    docx_path = material_root / "scripts" / "sample.docx"
    doc = Document()
    doc.add_paragraph("DOCX短剧素材：旧城早餐摊遇到债务危机。")
    doc.add_paragraph("第二段强调付费墙反转和家庭误会。")
    doc.save(docx_path)

    pdf_path = material_root / "scripts" / "sample.pdf"
    pdf_path.write_bytes(b"%PDF fake body for monkeypatched extractor")

    def fake_pdf_extract(path):
        assert path == pdf_path
        return {
            "content": "PDF短剧素材：外卖骑手发现隐藏账本，结尾反转。",
            "parser": "pypdf",
            "truncated": False,
            "metadata": {"page_count": 1, "pages_indexed": 1},
        }

    monkeypatch.setattr(full_text_search_service, "_extract_pdf_text", fake_pdf_extract)

    db = _make_search_session()
    try:
        rebuild = full_text_search_service.rebuild_search_index(db)

        assert rebuild["available"] is True
        assert rebuild["indexed_count"] == 2
        assert rebuild["doc_type_counts"] == {"material": 2}

        result = full_text_search_service.search_documents(db, "短剧素材", doc_types=["material"], limit=10)
        assert result["count"] == 2
        extensions = {item["metadata"]["extension"] for item in result["items"]}
        parsers = {item["metadata"]["parser"] for item in result["items"]}
        assert extensions == {".docx", ".pdf"}
        assert parsers == {"python-docx", "pypdf"}
        assert all(item["metadata"]["content_chars"] > 0 for item in result["items"])
        assert all(item["metadata"]["truncated"] is False for item in result["items"])
        pdf_item = next(item for item in result["items"] if item["metadata"]["extension"] == ".pdf")
        docx_item = next(item for item in result["items"] if item["metadata"]["extension"] == ".docx")
        assert pdf_item["metadata"]["requires_rights_review"] is True
        assert "pdf_material" in pdf_item["metadata"]["policy_flags"]
        assert docx_item["metadata"]["requires_rights_review"] is False
    finally:
        db.close()


def test_full_text_search_skips_oversized_binary_material_files(monkeypatch, tmp_path):
    from app.services import full_text_search_service

    _, material_root = _isolate_file_libraries(monkeypatch, tmp_path)
    (material_root / "large.docx").write_bytes(b"oversized")
    monkeypatch.setattr(full_text_search_service, "MAX_INDEXABLE_FILE_BYTES", 4)

    db = _make_search_session()
    try:
        rebuild = full_text_search_service.rebuild_search_index(db)

        assert rebuild["available"] is True
        assert rebuild["indexed_count"] == 0
        assert rebuild["doc_type_counts"] == {}
    finally:
        db.close()


def test_material_ingestion_audit_marks_image_ocr_boundary(monkeypatch, tmp_path):
    from app.services import full_text_search_service

    _, material_root = _isolate_file_libraries(monkeypatch, tmp_path)
    (material_root / "notes.txt").write_text("普通文本素材", encoding="utf-8")
    (material_root / "scan.png").write_bytes(
        b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
        b"\x08\x02\x00\x00\x00\x90wS\xde\x00\x00\x00\x0cIDATx\x9cc```\x00"
        b"\x00\x00\x04\x00\x01\xf6\x178U\x00\x00\x00\x00IEND\xaeB`\x82"
    )
    monkeypatch.delenv(full_text_search_service.MATERIAL_OCR_ENV, raising=False)

    audit = full_text_search_service.get_material_ingestion_audit()

    assert audit["available"] is True
    assert audit["supported_files"] == 2
    assert audit["ocr_enabled"] is False
    assert audit["ocr_candidate_files"] == 1
    assert audit["review_required_files"] == 1
    assert audit["status_counts"]["indexable"] == 1
    assert audit["status_counts"]["image_ocr_disabled"] == 1


def test_image_material_ocr_is_opt_in_and_marks_rights_review(monkeypatch, tmp_path):
    from PIL import Image

    from app.services import full_text_search_service

    _, material_root = _isolate_file_libraries(monkeypatch, tmp_path)
    image_path = material_root / "scan.png"
    Image.new("RGB", (16, 16), color="white").save(image_path)
    monkeypatch.setenv(full_text_search_service.MATERIAL_OCR_ENV, "1")
    monkeypatch.setattr(full_text_search_service, "_extract_text_from_image", lambda _image: "OCR短剧素材：截图里出现旧城追债反转。")

    db = _make_search_session()
    try:
        rebuild = full_text_search_service.rebuild_search_index(db)

        assert rebuild["available"] is True
        assert rebuild["indexed_count"] == 1
        result = full_text_search_service.search_documents(db, "截图", doc_types=["material"], limit=10)
        assert result["count"] == 1
        metadata = result["items"][0]["metadata"]
        assert metadata["parser"] == "tesseract_image"
        assert metadata["ocr_used"] is True
        assert metadata["requires_rights_review"] is True
        assert metadata["allowed_for_direct_copy"] is False
        assert "copyright_review_required" in metadata["policy_flags"]
    finally:
        db.close()


def test_scanned_pdf_ocr_fallback_is_opt_in(monkeypatch, tmp_path):
    import pypdf

    from app.services import full_text_search_service

    _, material_root = _isolate_file_libraries(monkeypatch, tmp_path)
    pdf_path = material_root / "scanned.pdf"
    pdf_path.write_bytes(b"%PDF fake scanned body")

    class FakePage:
        def extract_text(self):
            return ""

    class FakeReader:
        is_encrypted = False
        pages = [FakePage()]

    monkeypatch.setattr(pypdf, "PdfReader", lambda _path: FakeReader())
    monkeypatch.setenv(full_text_search_service.MATERIAL_OCR_ENV, "1")
    monkeypatch.setattr(
        full_text_search_service,
        "_extract_pdf_ocr_text",
        lambda path, page_count: {
            "content": "扫描PDF短剧素材：合同截图暴露隐藏股权。",
            "parser": "tesseract_pdf",
            "truncated": False,
            "metadata": {
                "page_count": page_count,
                "pages_indexed": 1,
                "ocr_used": True,
                "ocr_source": "scanned_pdf",
            },
        },
    )

    db = _make_search_session()
    try:
        rebuild = full_text_search_service.rebuild_search_index(db)

        assert rebuild["available"] is True
        assert rebuild["indexed_count"] == 1
        result = full_text_search_service.search_documents(db, "隐藏股权", doc_types=["material"], limit=10)
        assert result["count"] == 1
        metadata = result["items"][0]["metadata"]
        assert metadata["parser"] == "tesseract_pdf"
        assert metadata["page_count"] == 1
        assert metadata["ocr_used"] is True
        assert metadata["requires_rights_review"] is True
    finally:
        db.close()


def test_full_text_search_indexes_skill_plugin_documents(monkeypatch, tmp_path):
    from app.services import full_text_search_service

    _isolate_file_libraries(monkeypatch, tmp_path)
    skill_root = full_text_search_service.SKILL_LIBRARY_ROOT
    (skill_root / "short-drama-writer").mkdir()
    (skill_root / "short-drama-writer" / "SKILL.md").write_text(
        "# 短剧写作技能\n\n用于爆款短剧大纲、角色关系和前三集样稿生成。",
        encoding="utf-8",
    )
    (skill_root / "integrated-skills").mkdir()
    (skill_root / "integrated-skills" / "manifest.json").write_text(
        '{"skills": ["short-drama-writer", "short-drama-reviewer"]}',
        encoding="utf-8",
    )
    (skill_root / "third_party").mkdir()
    (skill_root / "third_party" / "README.md").write_text(
        "第三方原始仓库不进入本地技能事实源。",
        encoding="utf-8",
    )

    db = _make_search_session()
    try:
        rebuild = full_text_search_service.rebuild_search_index(db)

        assert rebuild["available"] is True
        assert rebuild["indexed_count"] == 2
        assert rebuild["doc_type_counts"] == {"skill": 2}

        result = full_text_search_service.search_documents(db, "短剧写作", doc_types=["skill"], limit=10)
        assert result["count"] == 1
        assert result["items"][0]["doc_type"] == "skill"
        assert result["items"][0]["metadata"]["kind"] == "skill_file"
        assert result["items"][0]["source_id"].startswith("file:05-skill-plugins:")
    finally:
        db.close()


def test_full_text_search_skips_unreadable_file_library_entries(monkeypatch, tmp_path):
    from app.services import full_text_search_service

    rule_root, material_root = _isolate_file_libraries(monkeypatch, tmp_path)
    (rule_root / "core.md").write_text("规则文件仍应进入索引。", encoding="utf-8")
    (material_root / "broken.json").write_bytes(b'{"bad": "\xff"}')

    db = _make_search_session()
    try:
        rebuild = full_text_search_service.rebuild_search_index(db)

        assert rebuild["available"] is True
        assert rebuild["indexed_count"] == 1
        assert rebuild["doc_type_counts"] == {"rule": 1}
    finally:
        db.close()
