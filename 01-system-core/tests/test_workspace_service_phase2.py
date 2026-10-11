from __future__ import annotations

from pathlib import Path
import importlib.util
import json
import sqlite3
import sys

from docx import Document
import pytest


def test_runtime_workspace_paths_do_not_hardcode_local_desktop():
    repo_root = Path(__file__).resolve().parents[1]
    runtime_files = [
        repo_root / "app" / "services" / "acceptance_artifact_service.py",
        repo_root / "app" / "routers" / "episode_segment_router.py",
    ]

    for runtime_file in runtime_files:
        text = runtime_file.read_text(encoding="utf-8")
        assert "/Users/" not in text


def test_acceptance_artifact_dir_is_project_relative():
    from app.services import workspace_service
    from app.services.acceptance_artifact_service import DESKTOP_ACCEPTANCE_DIR

    assert DESKTOP_ACCEPTANCE_DIR == workspace_service.PROJECT_LIBRARY_ROOT / "墨戏生花验收"


def test_story_source_docx_is_imported_to_controlled_library(monkeypatch, tmp_path):
    from app.services import workspace_service
    from app.services.story_source_service import load_story_source

    project_library = tmp_path / "03-project-library"
    monkeypatch.setattr(workspace_service, "PROJECT_LIBRARY_ROOT", project_library)
    external_dir = tmp_path / "external"
    external_dir.mkdir()
    source_doc = external_dir / "我的原稿.docx"

    document = Document()
    document.add_paragraph("测试短剧")
    document.add_paragraph("题材类型：校园，逆袭")
    document.add_paragraph("大纲：男主进入新学校后完成翻身。")
    document.save(source_doc)

    payload = load_story_source(source_doc)
    imported = Path(payload["source_doc_path"])

    assert imported.is_file()
    assert imported != source_doc.resolve()
    assert imported.is_relative_to(project_library / "imports" / "story-sources")
    assert payload["source_doc_name"] == "我的原稿.docx"
    assert payload["title"] == "测试短剧"


def test_activity_card_raw_body_does_not_use_v_html():
    core_root = Path(__file__).resolve().parents[1]
    activity_card = core_root / "webui" / "src" / "components" / "feed" / "ActivityCard.vue"
    text = activity_card.read_text(encoding="utf-8")

    assert "v-html" not in text
    assert "rawBodyTokens" in text


def test_generation_job_uses_thread_for_sync_model_call():
    core_root = Path(__file__).resolve().parents[1]
    generation_routes = core_root / "app" / "routers" / "projects" / "routes_generation.py"
    text = generation_routes.read_text(encoding="utf-8")

    assert "await asyncio.to_thread(" in text
    assert "generate_text," in text


def _load_migration_module():
    core_root = Path(__file__).resolve().parents[1]
    script_path = core_root / "scripts" / "migrate_legacy_workspaces.py"
    spec = importlib.util.spec_from_file_location("migrate_legacy_workspaces", script_path)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def test_legacy_workspace_migration_dry_run_does_not_move(tmp_path):
    module = _load_migration_module()
    project_library = tmp_path / "03-project-library"
    legacy_project = project_library / "新建项目" / "旧项目"
    legacy_project.mkdir(parents=True)
    (legacy_project / "project.json").write_text("{}", encoding="utf-8")

    report = module.migrate_legacy_workspaces(project_library_root=project_library)

    assert report.dry_run is True
    assert report.items[0].status == "would_move"
    assert legacy_project.exists()
    assert not (project_library / "new-projects" / "旧项目").exists()


def test_legacy_workspace_migration_apply_moves_and_merges_index(tmp_path):
    module = _load_migration_module()
    project_library = tmp_path / "03-project-library"
    legacy_root = project_library / "新建项目"
    legacy_project = legacy_root / "旧项目"
    legacy_project.mkdir(parents=True)
    (legacy_project / "project.json").write_text('{"name":"旧项目"}', encoding="utf-8")
    (legacy_root / "_index.json").write_text(
        '{"projects":[{"id":"p1","name":"旧项目","path":"' + str(legacy_project) + '"}]}',
        encoding="utf-8",
    )

    report = module.migrate_legacy_workspaces(project_library_root=project_library, apply=True)
    new_project = project_library / "new-projects" / "旧项目"

    assert report.dry_run is False
    assert report.items[0].status == "moved"
    assert not legacy_project.exists()
    assert (new_project / "project.json").is_file()
    index_text = (project_library / "new-projects" / "_index.json").read_text(encoding="utf-8")
    assert str(new_project) in index_text


def test_legacy_workspace_migration_can_update_project_v2_db(tmp_path):
    module = _load_migration_module()
    project_library = tmp_path / "03-project-library"
    legacy_project = project_library / "新建项目" / "旧项目"
    legacy_project.mkdir(parents=True)
    new_project = project_library / "new-projects" / "旧项目"
    db_path = tmp_path / "tasks.db"

    with sqlite3.connect(db_path) as conn:
        conn.execute("CREATE TABLE project_v2 (id TEXT PRIMARY KEY, workspace_path TEXT)")
        conn.execute("INSERT INTO project_v2 (id, workspace_path) VALUES (?, ?)", ("p1", str(legacy_project)))
        conn.commit()

    report = module.migrate_legacy_workspaces(
        project_library_root=project_library,
        apply=True,
        update_db=True,
        db_path=db_path,
    )

    with sqlite3.connect(db_path) as conn:
        saved_path = conn.execute("SELECT workspace_path FROM project_v2 WHERE id = 'p1'").fetchone()[0]

    assert report.db_updated == 1
    assert saved_path == str(new_project)


def test_short_drama_export_uses_workspace_service_artifact_root(monkeypatch, tmp_path):
    from app.services import workspace_service
    from app.services.short_drama_export import export_txt

    project_library = tmp_path / "03-project-library"
    monkeypatch.setattr(workspace_service, "PROJECT_LIBRARY_ROOT", project_library)
    metadata = {
        "project": {"name": "导出项目"},
        "episodes": {
            "1": {
                "steps": {
                    "script_draft": {
                        "full_output": "【场1-1：日 内 教室】\n△男主推门。\n男主：我来了。"
                    }
                }
            }
        },
    }

    path = export_txt("导出项目", metadata, 1)

    assert path.is_file()
    assert path.is_relative_to(project_library / "墨戏生花验收")
    assert "男主：我来了。" in path.read_text(encoding="utf-8")


def test_short_drama_export_rejects_empty_script_range(monkeypatch, tmp_path):
    from app.services import workspace_service
    from app.services.short_drama_export import NoExportableScriptError, export_txt

    project_library = tmp_path / "03-project-library"
    monkeypatch.setattr(workspace_service, "PROJECT_LIBRARY_ROOT", project_library)
    metadata = {
        "project": {"name": "空导出项目", "episode_count": 3},
        "episodes": {
            "1": {"steps": {"concept": {"full_output": "只有概念，没有正文"}}},
        },
    }

    with pytest.raises(NoExportableScriptError, match="未找到可导出的剧本正文"):
        export_txt("空导出项目", metadata, 1, scope="all_episodes")

    export_root = project_library / "墨戏生花验收"
    assert not export_root.exists()


def test_workspace_service_centralizes_episode_scene_checkpoint_paths(tmp_path):
    from app.services import workspace_service

    workspace = tmp_path / "03-project-library" / "new-projects" / "项目A"
    workspace.mkdir(parents=True)

    episode_dir = workspace_service.episode_dir(workspace, 3, create=True)
    scene_file = workspace_service.scene_version_path(
        workspace,
        episode_number=3,
        scene_number=2,
        version_code="v001",
    )
    checkpoint_file = workspace_service.checkpoint_path(workspace, "cp001")

    assert episode_dir == workspace / "episodes" / "ep03"
    assert scene_file == workspace / "episodes" / "ep03" / "scenes" / "sc02_v001.md"
    assert checkpoint_file == workspace / "checkpoints" / "cp001.json"
    assert episode_dir.is_dir()
    assert scene_file.parent.is_dir()
    assert checkpoint_file.parent.is_dir()


@pytest.mark.asyncio
async def test_resume_testing_workflow_reads_controlled_workspace(monkeypatch, tmp_path):
    from app.routers.short_drama.resume_routes import resume_testing_workflow
    from app.schemas.short_drama_writing import TestingResumeQuery
    from app.services import workspace_service

    project_library = tmp_path / "03-project-library"
    workspace_root = project_library / "new-projects"
    monkeypatch.setattr(workspace_service, "PROJECT_LIBRARY_ROOT", project_library)
    monkeypatch.setattr(workspace_service, "WORKSPACE_ROOT", workspace_root)
    monkeypatch.setattr(workspace_service, "LEGACY_WORKSPACE_ROOT", project_library / "新建项目")

    workspace = workspace_root / "恢复项目"
    workspace.mkdir(parents=True)
    (workspace / "frontend_workflow_progress.json").write_text(
        json.dumps(
            {
                "project_id": "p1",
                "project_name": "恢复项目",
                "updated_at": "2026-05-09T20:00:00",
                "results": [
                    {"episode_no": 1, "char_count": 1200, "blocking_issues": []},
                    {"episode_no": 2, "char_count": 0, "blocking_issues": ["empty"]},
                ],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    response = await resume_testing_workflow(
        TestingResumeQuery(workspace_path=str(workspace), expected_total=3),
        _=True,
    )

    assert response.workspace_path == str(workspace.resolve())
    assert response.project_id == "p1"
    assert response.completed_episodes == [1]
    assert response.failed_episodes == [2]
    assert response.pending_episodes == [2, 3]


@pytest.mark.asyncio
async def test_resume_testing_workflow_rejects_out_of_bounds_workspace(monkeypatch, tmp_path):
    from fastapi import HTTPException

    from app.routers.short_drama.resume_routes import resume_testing_workflow
    from app.schemas.short_drama_writing import TestingResumeQuery
    from app.services import workspace_service

    project_library = tmp_path / "03-project-library"
    monkeypatch.setattr(workspace_service, "PROJECT_LIBRARY_ROOT", project_library)
    monkeypatch.setattr(workspace_service, "WORKSPACE_ROOT", project_library / "new-projects")
    monkeypatch.setattr(workspace_service, "LEGACY_WORKSPACE_ROOT", project_library / "新建项目")
    outside = tmp_path / "outside-workspace"
    outside.mkdir()
    (outside / "frontend_workflow_progress.json").write_text("{}", encoding="utf-8")

    with pytest.raises(HTTPException) as exc:
        await resume_testing_workflow(
            TestingResumeQuery(workspace_path=str(outside), expected_total=1),
            _=True,
        )

    assert exc.value.status_code == 400


def test_project_routes_no_longer_build_runtime_workspace_files_directly():
    core_root = Path(__file__).resolve().parents[1]
    files = [
        core_root / "app" / "routers" / "projects" / "routes_episodes_scenes.py",
        core_root / "app" / "routers" / "projects" / "routes_generation.py",
    ]

    for file_path in files:
        text = file_path.read_text(encoding="utf-8")
        assert "Path(project.workspace_path)" not in text
        assert "/ \"episodes\" /" not in text
        assert "/ \"generation_jobs\" /" not in text
